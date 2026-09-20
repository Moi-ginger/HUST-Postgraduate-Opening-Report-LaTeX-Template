from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
SOURCE = Path("/Users/GillianChen/Documents/Cursor Project/MyPaper/陈思静-开题报告-sep20.docx")
CARDS_SOURCE = Path("/Users/GillianChen/Documents/Cursor Project/MyPaper/陈思静-文献阅读卡-40篇-V1核校版.docx")
OUT_REVIEW = ROOT / "body" / "literature-review.tex"
OUT_SELECTION = ROOT / "body" / "selection-report.tex"
OUT_CARDS = ROOT / "body" / "literature-cards.tex"
OUT_BIB = ROOT / "ref" / "report.bib"


def text_of_paragraph(element):
    return "".join(node.text or "" for node in element.iter() if node.tag == qn("w:t")).strip()


def tex_escape(value):
    value = value.replace("\u200b", "").replace("\ufeff", "")
    value = value.replace("\\", r"\textbackslash{}")
    replacements = {
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_",
        "{": r"\{",
        "}": r"\}",
        "~": r"\textasciitilde{}",
        "^": r"\textasciicircum{}",
    }
    for old, new in replacements.items():
        value = value.replace(old, new)
    # Turn source-document numeric markers such as [3][20] into real
    # biblatex citations. The bibliography keys are stable across rebuilds.
    value = re.sub(
        r"(?:\[(\d+)\])+",
        lambda m: r"\cite{" + ",".join(f"ref{int(n):02d}" for n in re.findall(r"\[(\d+)\]", m.group(0))) + "}",
        value,
    )
    return value


def heading_command(value, chapter_kind):
    if chapter_kind == "review":
        match = re.match(r"^1\.(\d+)\.(\d+)\s+(.+)$", value)
        if match:
            return rf"\subsection{{{tex_escape(match.group(3))}}}"
        match = re.match(r"^1\.(\d+)\s+(.+)$", value)
        if match:
            return rf"\section{{{tex_escape(match.group(2))}}}"
    else:
        match = re.match(r"^(\d+)\.\s+(.+)$", value)
        if match:
            return rf"\section{{{tex_escape(match.group(2))}}}"
        match = re.match(r"^(\d+)\.(\d+)\.\s+(.+)$", value)
        if match:
            return rf"\subsection{{{tex_escape(match.group(3))}}}"
        match = re.match(r"^(\d+)\.(\d+)\.(\d+)\.\s+(.+)$", value)
        if match:
            return rf"\subsubsection{{{tex_escape(match.group(4))}}}"
    return None


def table_tex(table, caption, number):
    rows = []
    for row in table.rows:
        cells = [" ".join(cell.text.split()) for cell in row.cells]
        rows.append(cells)
    if not rows:
        return ""
    ncols = len(rows[0])
    if ncols == 6:
        spec = r">{\raggedright\arraybackslash}p{1.85cm} *{5}{>{\raggedright\arraybackslash}X}"
    elif ncols == 5:
        spec = r">{\raggedright\arraybackslash}p{2.0cm} *{4}{>{\raggedright\arraybackslash}X}"
    elif ncols == 4:
        spec = r">{\raggedright\arraybackslash}p{2.6cm} *{3}{>{\raggedright\arraybackslash}X}"
    else:
        spec = r">{\raggedright\arraybackslash}p{2.7cm} *{2}{>{\raggedright\arraybackslash}X}"
    lines = [
        r"\begin{table}[htbp]",
        r"\centering",
        rf"\caption{{{tex_escape(caption)}}}",
        rf"\label{{tab:report-{number}}}",
        r"\scriptsize",
        r"\renewcommand{\arraystretch}{1.18}",
        rf"\begin{{tabularx}}{{\textwidth}}{{{spec}}}",
        r"\toprule",
    ]
    header = rows[0]
    lines.append(" & ".join(r"\textbf{" + tex_escape(c) + "}" for c in header) + r" \\")
    lines.append(r"\midrule")
    for row in rows[1:]:
        if len(row) < ncols:
            row += [""] * (ncols - len(row))
        lines.append(" & ".join(tex_escape(c) for c in row[:ncols]) + r" \\")
    lines += [r"\bottomrule", r"\end{tabularx}", r"\end{table}", ""]
    return "\n".join(lines)


def card_field(text, label):
    prefix = label + " "
    if text.startswith(prefix):
        return text[len(prefix):].strip()
    return text[len(label):].strip() if text.startswith(label) else text


def render_cards(doc):
    starts = [
        i for i, paragraph in enumerate(doc.paragraphs)
        if paragraph.text.strip() == "华中科技大学设计学院硕士研究生"
    ]
    output = [
        r"\newgeometry{left=1.25in,right=1.25in,top=1in,bottom=1in}",
        r"\phantomsection",
        r"\chapter*{一、文献阅读卡}\label{chap:literature-cards}",
        r"\addcontentsline{toc}{chapter}{文献阅读卡}",
        r"\clearpage",
        r"\setlength{\parindent}{0pt}",
        r"\setlength{\parskip}{0pt}",
        r"\setstretch{1.5}",
        r"",
    ]
    for card_no, start in enumerate(starts, 1):
        end = starts[card_no] if card_no < len(starts) else len(doc.paragraphs)
        paragraphs = [p.text.strip() for p in doc.paragraphs[start:end] if p.text.strip()]
        table = doc.tables[card_no - 1]
        metadata = [cell.text.strip() for cell in table.rows[0].cells]
        student_id = metadata[0].replace("学号：", "").strip()
        name = metadata[1].replace("姓名：", "").strip()
        serial = metadata[2].replace("序号：", "").strip()
        title = card_field(paragraphs[2], "名称")
        authors = card_field(paragraphs[3], "作者")
        source = card_field(paragraphs[4], "出处")

        if card_no > 1:
            output.append(r"\clearpage")
        output += [
            r"\begin{center}",
            r"{\songti\bfseries\fontsize{18.5pt}{27.75pt}\selectfont 华中科技大学设计学院硕士研究生\par}",
            r"{\songti\bfseries\fontsize{18.5pt}{27.75pt}\selectfont 文献阅读卡\par}",
            r"\end{center}",
            r"\vspace{0.25em}",
            rf"\noindent\underline{{\makebox[\textwidth][l]{{\textbf{{学号：{tex_escape(student_id)}\hfill 姓名：{tex_escape(name)}\hfill 序号：{tex_escape(serial)}}}}}}}\par",
            r"\vspace{0.35em}",
            rf"\noindent\textbf{{名称}}\ {tex_escape(title)}\par",
            rf"\noindent\textbf{{作者}}\ {tex_escape(authors)}\quad 译者\par",
            rf"\noindent\textbf{{出处}}\ {tex_escape(source)}\quad 页\par",
            r"\vspace{0.35em}",
            r"\noindent\textbf{文献阅读心得（每篇不少于150字）：}\par",
        ]
        body_paragraphs = paragraphs[6:]
        body_index = 0
        while body_index < len(body_paragraphs):
            paragraph = body_paragraphs[body_index]
            if paragraph in ("文献阅读心得：", "文献阅读心得（每篇不少于150字）："):
                body_index += 1
                continue
            labels = ["研究背景", "研究问题", "研究方法", "主要发现", "贡献与意义"]
            found = next((label for label in labels if paragraph.startswith(label)), None)
            if found:
                rest = paragraph[len(found):].strip()
                if not rest:
                    if body_index + 1 < len(body_paragraphs) and body_paragraphs[body_index + 1].startswith("RQ1"):
                        output.append(rf"\noindent\textbf{{{found}}}\ {tex_escape(body_paragraphs[body_index + 1])}\par")
                        body_index += 1
                    else:
                        output.append(rf"\noindent\textbf{{{found}}}\par")
                else:
                    output.append(rf"\noindent\textbf{{{found}}}\ {tex_escape(rest)}\par")
            else:
                output.append(tex_escape(paragraph) + r"\par")
            body_index += 1
    output += [r"\restoregeometry", ""]
    return "\n".join(output)


def build_events(doc):
    events = []
    paragraph_index = 0
    table_index = 0
    for child in doc._element.body.iterchildren():
        if child.tag == qn("w:p"):
            events.append(("p", paragraph_index, text_of_paragraph(child)))
            paragraph_index += 1
        elif child.tag == qn("w:tbl"):
            events.append(("t", table_index, doc.tables[table_index]))
            table_index += 1
    return events


def render_review(doc, events):
    refs = [p.text.strip() for p in doc.paragraphs[774:814] if p.text.strip()]
    output = [r"\chapter{文献综述}\label{chap:literature-review}", ""]
    pending_caption = None
    for kind, index, value in events:
        if kind == "p":
            if index < 702 or index > 772:
                continue
            if value.startswith("表") and re.match(r"^表\d+\s+", value):
                pending_caption = re.sub(r"^表\d+\s+", "", value)
                continue
            command = heading_command(value, "review")
            if command:
                output += [command, ""]
            elif value:
                output += [tex_escape(value), ""]
        else:
            if index in (0, 1, 2):
                caption = pending_caption or f"文献综述比较表{index + 1}"
                output.append(table_tex(value, caption, index + 1))
                pending_caption = None
    return "\n".join(output), refs


def render_selection(events):
    output = [r"\chapter{选题报告}\label{chap:selection-report}", ""]
    pending_caption = None
    for kind, index, value in events:
        if kind == "p":
            if not (828 <= index <= 951):
                continue
            if 841 <= index <= 912:
                continue
            if value == "6. 参考文献":
                continue
            if value.startswith("表") and re.match(r"^表\d+\s+", value):
                pending_caption = re.sub(r"^表\d+\s+", "", value)
                continue
            command = heading_command(value, "selection")
            if command:
                output += [command, ""]
                if value.startswith("4. "):
                    output[-1:-1] = [
                        r"\begin{figure}[htbp]",
                        r"\centering",
                        r"\includegraphics[width=0.92\textwidth]{source-image2.png}",
                        r"\caption{研究框架与实施路径}",
                        r"\label{fig:research-framework}",
                        r"\end{figure}",
                        "",
                    ]
            elif value.startswith("图1 "):
                continue
            elif value:
                output += [tex_escape(value), ""]
        else:
            if index in (6, 7, 8):
                fallback_captions = {
                    6: "实验条件设置",
                    7: "实验测量维度与指标",
                    8: "研究计划与进度安排",
                }
                caption = pending_caption or fallback_captions[index]
                output.append(table_tex(value, caption, index + 1))
                pending_caption = None
    # The original report numbers the next block as section 3 because section 2
    # is the literature review. It is kept as a cross-reference without copying
    # the same paragraphs a second time.
    marker = r"\section{国内外研究综述}"
    if marker not in output:
        # Insert the reference section immediately after the first section's body.
        for pos, item in enumerate(output):
            if item == r"\section{研究问题与方法}":
                output[pos:pos] = [marker, r"本报告的国内外研究综述见\hyperref[chap:literature-review]{第\ref{chap:literature-review}章“文献综述”}，此处不再重复排版。", ""]
                break
    return "\n".join(output)


def bib_escape(value):
    return (value.replace("\\", "\\textbackslash{}")
                 .replace("{", "\\{")
                 .replace("}", "\\}")
                 .replace("%", "\\%")
                 .replace("#", "\\#")
                 .replace("&", "\\&"))


def bib_authors(raw):
    """Convert the compact GB/T-style author string to BibLaTeX names."""
    raw = re.sub(r",?\s*et al\.?\s*$", "", raw, flags=re.IGNORECASE)
    names = []
    for item in raw.split(","):
        item = item.strip()
        if not item:
            continue
        match = re.match(r"^(.*?)\s+((?:[A-ZÁÉÍÓÚÜÑ](?:\s+|$))+)\s*$", item)
        if match:
            surname = match.group(1).strip()
            initials = " ".join(match.group(2).split())
            names.append(f"{surname}, {initials}")
        else:
            names.append(item)
    return " and ".join(names) if names else "Unknown"


def parse_reference(raw):
    """Parse the source report's GB/T-style reference into BibLaTeX fields."""
    match = re.match(
        r"^(?P<authors>.+?)\.\s+(?P<title>.*?)\[(?P<kind>J|C|M|R|EB/OL)\]"
        r"(?://(?P<container>.*?))?\.\s+(?P<tail>.*)$",
        raw,
    )
    if not match:
        return {"entrytype": "misc", "author": "Unknown", "title": raw, "year": "2026"}

    data = match.groupdict()
    tail = data["tail"].strip()
    year_match = re.search(r"(?<!\d)((?:19|20)\d{2})(?:[-–](?:19|20)\d{2})?", tail)
    year = year_match.group(1) if year_match else "2026"
    before_year = tail[:year_match.start()].strip(" .,") if year_match else tail
    after_year = tail[year_match.end():].strip(" .,") if year_match else ""
    page_match = re.search(r":\s*([0-9][0-9:–-]*)\s*\.?$", tail)
    pages = page_match.group(1).replace("–", "--") if page_match else None

    kind = data["kind"]
    if kind == "J":
        entrytype, fields = "article", {"journaltitle": before_year}
    elif kind == "C":
        entrytype, fields = "inproceedings", {"booktitle": (data["container"] or before_year).strip()}
    elif kind == "M":
        entrytype, fields = "inbook", {"booktitle": (data["container"] or before_year).strip()}
    elif kind == "R":
        entrytype, fields = "techreport", {"institution": before_year}
    else:
        entrytype, fields = "online", {"note": tail}

    fields.update({
        "author": bib_authors(data["authors"]),
        "title": data["title"].strip(),
        "year": year,
    })
    if pages:
        fields["pages"] = pages

    volume_match = re.search(r"(?:^|,\s*)(\d+)(?:\(([^)]+)\))?", after_year)
    if volume_match and kind in {"J", "C"}:
        fields["volume"] = volume_match.group(1)
        if volume_match.group(2):
            fields["number"] = volume_match.group(2)

    doi_match = re.search(r"DOI:\s*([^,.\s]+)", tail, flags=re.IGNORECASE)
    if doi_match:
        fields["doi"] = doi_match.group(1)

    return {"entrytype": entrytype, **fields}


def refs_bib(refs):
    lines = ["% Generated from the reference list in the source report.", "% Numeric keys are kept stable so citations remain reproducible.", ""]
    for ref in refs:
        match = re.match(r"^\[(\d+)\]\s*(.*)$", ref)
        if not match:
            continue
        number, raw = match.groups()
        parsed = parse_reference(raw)
        lines += [
            f"@{parsed.pop('entrytype')}{{ref{int(number):02d},",
        ]
        for field, value in parsed.items():
            lines.append(f"  {field} = {{{bib_escape(value)}}},")
        lines += ["}", ""]
    return "\n".join(lines)


def main():
    doc = Document(SOURCE)
    cards_doc = Document(CARDS_SOURCE)
    events = build_events(doc)
    review, refs = render_review(doc, events)
    selection = render_selection(events)
    cards = render_cards(cards_doc)
    OUT_REVIEW.write_text(review + "\n", encoding="utf-8")
    OUT_SELECTION.write_text(selection + "\n", encoding="utf-8")
    OUT_CARDS.write_text(cards + "\n", encoding="utf-8")
    OUT_BIB.write_text(refs_bib(refs), encoding="utf-8")
    print(f"wrote {OUT_REVIEW}")
    print(f"wrote {OUT_SELECTION}")
    print(f"wrote {OUT_CARDS}")
    print(f"wrote {OUT_BIB}")


if __name__ == "__main__":
    main()
